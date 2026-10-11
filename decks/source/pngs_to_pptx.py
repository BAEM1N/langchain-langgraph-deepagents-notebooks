#!/usr/bin/env python3
"""pngs_to_pptx.py — 슬라이드 PNG들을 16:9 PPTX에 풀블리드로 임베드.

사용: python tools/pngs_to_pptx.py <png_glob> <out.pptx>
  예: python tools/pngs_to_pptx.py "강의교안/01_beginner/_render/hp-*.png" 산출물/01_beginner.pptx
PNG는 파일명 정렬 순서대로 슬라이드가 된다. (HTML 덱 → PDF → PNG 캡처 후 사용)
"""
import glob
import json
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

def main():
    if len(sys.argv) < 3:
        print(__doc__); sys.exit(1)
    pngs = sorted(glob.glob(sys.argv[1]))
    if not pngs:
        print("PNG 없음:", sys.argv[1]); sys.exit(1)
    # 선택: 3번째 인자 = speeches.json (슬라이드 순서 대본) → 발표자 노트 주입
    speeches = []
    if len(sys.argv) > 3 and Path(sys.argv[3]).exists():
        speeches = json.loads(Path(sys.argv[3]).read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for i, png in enumerate(pngs):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)
        if i < len(speeches) and speeches[i]:
            slide.notes_slide.notes_text_frame.text = speeches[i]
    n_notes = sum(1 for s in speeches if s)
    prs.save(sys.argv[2])
    print(f"✅ {sys.argv[2]} — {len(pngs)}슬라이드 (풀블리드)" + (f", 노트 {n_notes}개" if speeches else ""))

if __name__ == "__main__":
    main()
